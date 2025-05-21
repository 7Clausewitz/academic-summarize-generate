from openai import OpenAI
import docx
import os
import re
import tkinter as tk
from tkinter import filedialog, scrolledtext, messagebox, ttk
import threading

client = OpenAI(
    base_url='https://xiaoai.plus/v1',
    # sk-xxx替换为自己的key
    api_key='sk-qtr0Y0KiEkwF8EH3mzP5uj0lXJUMqK1oEYnYPjm4hvMIW1Nx'
)

def read_docx(file_path):
    """读取指定的doc文件内容"""
    if not os.path.exists(file_path):
        return "文件不存在"
    
    doc = docx.Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return "\n".join(full_text)

def analyze_document(doc_content):
    """分析综述文档中的语言和逻辑不足"""
    prompt = f"""
你的任务是分析以下综述文档中存在的语言和逻辑不足之处，并提出具体的改进建议。请仔细阅读文档，并从以下几个方面进行评估：
1. 语言表达是否清晰、准确、专业
2. 逻辑结构是否连贯，论证是否有力
3. 内容的完整性和一致性

在分析时将存在问题的文本一一列举指出，并分析给出修改建议
    
    综述内容：
    {doc_content}
    """

    completion = client.chat.completions.create(
        model="gemini-2.0-pro-exp",
        messages=[
            {"role": "user", "content": prompt}
        ]
    )
    response = completion.choices[0].message.content
    response = re.sub(r'<think>.*?</think>', '', response, flags=re.DOTALL)
    
    return response

class SummaryAnalyzerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("综述文档分析器")
        self.root.geometry("800x600")
        self.root.minsize(700, 500)
        
        self.create_widgets()
        self.file_path = ""
        
    def create_widgets(self):
        # 创建主框架
        main_frame = tk.Frame(self.root, padx=10, pady=10)
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # 文件选择区域
        file_frame = tk.Frame(main_frame)
        file_frame.pack(fill=tk.X, pady=(0, 10))
        
        tk.Label(file_frame, text="文档路径:").pack(side=tk.LEFT, padx=(0, 5))
        
        self.file_entry = tk.Entry(file_frame)
        self.file_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))
        
        browse_btn = tk.Button(file_frame, text="浏览...", command=self.browse_file)
        browse_btn.pack(side=tk.LEFT)
        
        # 操作按钮区域
        btn_frame = tk.Frame(main_frame)
        btn_frame.pack(fill=tk.X, pady=(0, 10))
        
        self.analyze_btn = tk.Button(btn_frame, text="分析文档", command=self.start_analysis, width=15)
        self.analyze_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        # 进度条
        self.progress_var = tk.DoubleVar()
        self.progress = ttk.Progressbar(btn_frame, variable=self.progress_var, length=200, mode="indeterminate")
        self.progress.pack(side=tk.LEFT, fill=tk.X, expand=True)
        
        # 结果显示区域
        result_frame = tk.Frame(main_frame)
        result_frame.pack(fill=tk.BOTH, expand=True)
        
        tk.Label(result_frame, text="分析结果:").pack(anchor=tk.W)
        
        # 创建带滚动条的文本区域
        text_frame = tk.Frame(result_frame)
        text_frame.pack(fill=tk.BOTH, expand=True)
        
        scroll_y = tk.Scrollbar(text_frame)
        scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.result_text = scrolledtext.ScrolledText(text_frame, wrap=tk.WORD, font=("Microsoft YaHei", 10))
        self.result_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        scroll_y.config(command=self.result_text.yview)
        self.result_text.config(yscrollcommand=scroll_y.set)
        
        # 状态栏
        self.status_var = tk.StringVar()
        self.status_var.set("准备就绪")
        status_bar = tk.Label(self.root, textvariable=self.status_var, bd=1, relief=tk.SUNKEN, anchor=tk.W)
        status_bar.pack(side=tk.BOTTOM, fill=tk.X)
    
    def browse_file(self):
        file_path = filedialog.askopenfilename(
            title="选择综述文档",
            filetypes=[("Word文档", "*.docx *.doc"), ("所有文件", "*.*")]
        )
        if file_path:
            self.file_path = file_path
            self.file_entry.delete(0, tk.END)
            self.file_entry.insert(0, file_path)
    
    def start_analysis(self):
        self.file_path = self.file_entry.get()
        if not self.file_path:
            messagebox.showwarning("警告", "请先选择一个文档文件")
            return
        
        if not os.path.exists(self.file_path):
            messagebox.showerror("错误", "文件不存在，请检查路径")
            return
        
        # 禁用分析按钮并显示进度条
        self.analyze_btn.config(state=tk.DISABLED)
        self.progress.start()
        self.status_var.set("正在分析文档，请稍候...")
        self.result_text.delete(1.0, tk.END)
        
        # 使用线程进行分析，避免界面卡顿
        threading.Thread(target=self.analyze_in_thread, daemon=True).start()
    
    def analyze_in_thread(self):
        try:
            # 读取文档内容
            doc_content = read_docx(self.file_path)
            if doc_content == "文件不存在":
                self.show_error("错误：指定的文件不存在，请检查路径是否正确。")
                return
            
            # 分析文档
            analysis_result = analyze_document(doc_content)
            
            # 在主线程中更新UI
            self.root.after(0, self.update_result, analysis_result)
            
        except Exception as e:
            self.root.after(0, self.show_error, f"分析过程中出错: {str(e)}")
    
    def update_result(self, result):
        self.result_text.delete(1.0, tk.END)
        self.result_text.insert(tk.END, result)
        self.progress.stop()
        self.analyze_btn.config(state=tk.NORMAL)
        self.status_var.set("分析完成")
    
    def show_error(self, error_msg):
        self.progress.stop()
        self.analyze_btn.config(state=tk.NORMAL)
        self.status_var.set("发生错误")
        messagebox.showerror("错误", error_msg)

def main():
    root = tk.Tk()
    app = SummaryAnalyzerApp(root)
    root.mainloop()

if __name__ == "__main__":
    main()